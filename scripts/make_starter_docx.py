"""Build the two starter Word files for Mystery Box: Word Challenge (report-rescue/).

Run from the repo root:  python3 scripts/make_starter_docx.py
The text must match TITLE, PARA, HEADING and FEATURES in report-rescue/data.js.
"""
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt

TITLE = 'My Dream Gaming Room'
PARA = ('My dream gaming room is the best place to play and relax with my freinds. The walls are dark blue, '
        'and bright lights glow purple and green. There is a soft sofa, a mini fridge full of snaks and a big '
        'shelf of games. It is an awsome room!')
HEADING = 'Room Features'
FEATURES = ['A huge curved screen for racing games', 'Two loud speakers for music', 'A gaming chair that glows in the dark']


def new_doc():
    d = Document()
    s = d.sections[0]
    s.orientation, s.page_width, s.page_height = WD_ORIENT.PORTRAIT, Mm(210), Mm(297)
    d.core_properties.title = TITLE
    return d


def plain(path):
    d = new_doc()
    for line in [TITLE, '', PARA, '', HEADING, *FEATURES]:
        d.add_paragraph(line)
    d.save(path)


def ready(path):
    d = new_doc()
    t = d.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(TITLE)
    r.bold, r.font.size = True, Pt(20)
    p = d.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.add_run(PARA).font.size = Pt(12)
    d.add_heading(HEADING, level=1)
    for f in FEATURES:
        b = d.add_paragraph(style='List Bullet')
        b.paragraph_format.line_spacing = 1.5
        b.add_run(f).font.size = Pt(12)
    d.save(path)


plain('report-rescue/dream-gaming-room.docx')
ready('report-rescue/dream-gaming-room-ready.docx')
print('done')
